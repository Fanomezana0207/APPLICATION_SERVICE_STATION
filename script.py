from pathlib import Path

import pandas

import openpyxl

from service import Service

if __name__ == "__main__":    

    CURENT_DIR = Path("__FILE__").parent.resolve()

    FILES_DIR = CURENT_DIR / "FILES" / "VENTE_PAR_CATEGORIE"

    SUIVI_SERVICE_FILE = Path("D:/FANOMEZANA/SUIVI SERVICE/OCTOBRE/RECAP SERVICE_OCTOBRE_2026.xlsx")

    #file_1 = "FILES/VENTE_PAR_CATEGORIE/category-sales-summary-2026-10-01-2026-10-01.xlsx"

    if not SUIVI_SERVICE_FILE.exists :
        raise FileNotFoundError(SUIVI_SERVICE_FILE)
    
    suivi_book = openpyxl.load_workbook(SUIVI_SERVICE_FILE)

    print(suivi_book.sheetnames)

    nom_feuille = "SERVICE ANOSIBE SEPTEMBRE "
    
    ws = suivi_book[nom_feuille]

    first_row = 4
    lavage_col = 3
    vidange_col = 6
    vulca_col = 9
    parallelisme_col = 12

    #ws["C4"] = 2000
    ws.cell(row=4,column=3).value=5000

    #suivi_book.save(SUIVI_SERVICE_FILE)

    #print("Modification enregisté")

    row = first_row
    for file in FILES_DIR.iterdir() :
        
        print("\n\t",file.stem[-10:],"\n")
        data = pandas.read_excel(file)
        list_service = []
        for i in range(len(data)):
            if i != 4 :
                serv = Service(data["CatÃ©gorie"][i],data["Ventes brutes"][1])
                list_service.append(serv)
        
        for service in list_service:
            if (service.getService() == "LAVAGE") :
                ws.cell(row=row,column=lavage_col).value = service.getVente()
            elif (service.getService() == "VIDANGE") :
                ws.cell(row=row,column=vidange_col).value = service.getVente()
            elif (service.getService() == "VULCA") :
                ws.cell(row=row,column=vulca_col).value = service.getVente()
            elif (service.getService() == "PARALLELISME & CARROSSAGE") :
                ws.cell(row=row,column=parallelisme_col).value = service.getVente()
            #print(service)
        
        row += 1

    suivi_book.save(SUIVI_SERVICE_FILE)
    
    print("Modification enregisté")

    #data = pandas.read_excel(file_1)
    # data = pandas.read_excel(file_1, sheet_name="category-sales-summary-2026-10-", usecols="A, C")

    #print(data)

    #list_service = []
    #for i in range(len(data)):
    #    if i != 4 :
    #        serv = Service(data["CatÃ©gorie"][i],data["Ventes brutes"][1])
    #        list_service.append(serv)

    #for service in list_service:
    #    print(service)

    #suivi_dates = pandas.read_excel(SUIVI_SERVICE_FILE, sheet_name="SERVICE ANOSIBE SEPTEMBRE ", usecols="A")

    



