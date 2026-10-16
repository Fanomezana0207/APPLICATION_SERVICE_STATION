from pathlib import Path

import pandas

import openpyxl

import datetime

from service import Service




file = Path(r"C:\Users\FEHIZORO\Downloads\sales-summary-2026-10-07-2026-10-13.csv")



SUIVI_SERVICE_FILE = Path("D:/FANOMEZANA/SUIVI SERVICE/OCTOBRE/RECAP SERVICE_OCTOBRE_2026.xlsx")
FIRST_ROW = 4
LAVAGE_COL = 3
VIDANGE_COL = 6
VULCA_COL = 9
PARALLELISME_COL = 12

def selectRowByDate (date_recherche, sheet) :
    for i in range(FIRST_ROW, sheet.max_row+1):
        d = sheet.cell(row = i, column = 1).value

        if d is None :
            continue

        if (isinstance(d,(datetime.datetime,datetime.date))) :
            d_en_str = d.strftime("%Y-%m-%d")
        else :
            d_en_str = str(d).strip()

        if d_en_str == date_recherche :
            return i
    print(f"date {date_recherche} non trouvé")
    return None

def insertServiceToExcel(service, sheet, row) :
    if (service.getService() == "LAVAGE") :
        sheet.cell(row=row,column=LAVAGE_COL).value = service.getVente()
    elif (service.getService() == "VIDANGE") :
        sheet.cell(row=row,column=VIDANGE_COL).value = service.getVente()
    elif (service.getService() == "VULCA") :
        sheet.cell(row=row,column=VULCA_COL).value = service.getVente()
    elif (service.getService() == "PARALLELISME & CARROSSAGE") :
        sheet.cell(row=row,column=PARALLELISME_COL).value = service.getVente()

def insertSalesCategories(dataFilePath, outputFilePath, sheetname = "") :
    dataFile = Path(dataFilePath)
    outputFile = Path(outputFilePath)
    data = pandas.read_csv(dataFile, encoding="utf-8-sig")
    data.columns = data.columns.str.strip() # Nettoyer les espaces éventuels

    if not outputFile.exists() :
        raise FileNotFoundError(outputFile)
        
    suivi_book = openpyxl.load_workbook(outputFile)
    if not(sheetname in suivi_book.sheetnames) :
        print(suivi_book.sheetnames)
        sheetname = input("veuiller COPIER-COLLER le nom de la feuille choisi : ")
    ws = suivi_book[sheetname]

    date_receipts = dataFile.stem[-10:]
    row = selectRowByDate(date_receipts,ws)

    for i in range(len(data)):
        serv = Service(data["Catégorie"][i],data["Ventes brutes"][i])
        insertServiceToExcel(serv, ws, row=row)

    suivi_book.save(outputFile)
    print(f"insertion des ventes par catégorie du {date_receipts}")
    print("Modification enregisté")

def insertAllSalesCategories(dataDirPath, outputFilePath) :
    dataDirPath = Path(dataDirPath)
    outputFile = Path(outputFilePath)
    suivi_book = openpyxl.load_workbook(outputFile)
    print(suivi_book.sheetnames)
    sheetname = input("veuiller COPIER-COLLER le nom de la feuille choisi : ")
    for file in dataDirPath.iterdir() :
        if file.is_file() and file.suffix.lower() == ".csv" :
            #dataFilePath = file.resolve()
            #print(dataFilePath)
            #print(type(dataFilePath))
            insertSalesCategories(file, outputFile,sheetname)

def generateReceipts(dataFilePath) :
    dataFilecsv = Path(dataFilePath)
    dataFileExcel = dataFilecsv.parent.resolve().joinpath(dataFilecsv.stem+".xlsx")
    data = pandas.read_csv(dataFilecsv, encoding="utf-8-sig")

    colonnes = [
        "Date",
        "Numéro du reçu",
        "Catégorie",
        "Article",
        "Ventes brutes",
        "Magasin",
        "Nom du client",
        "Statut"
    ]

    data = data[colonnes].copy()
    data["Ventes brutes"] = pandas.to_numeric(data["Ventes brutes"],errors="raise").astype("int64")
    #for i in range(1, len(data+1)) :
    #    if data["Nom du client"][i] != "" :
    #        data["Nom du client"][i].upper()

    data.to_excel(dataFileExcel, index=False)
    print("fichier receipts recalculé")

#dataFileToInput = input("Le lien du fichier sale-categorie à importer : ")
#dataFileOutput = input("Le lien du fichier Recap_service : ")
#insertSalesCategories(dataFileToInput, dataFileOutput)

#fileDir = input("Le lien du dossier où se trouve les fichiers : ").strip().strip("\"")
#dataFileOutput = input("Le lien du fichier Recap_service : ").strip().strip("\"")
#insertAllSalesCategories(fileDir, dataFileOutput)

#dataFilePath = input("Le lien du fichier receipts à importer : ").strip().strip("\"")
#generateReceipts(dataFilePath)